#!/usr/bin/env python3
"""Consolidate the 32-XPU tiered-prefix-cache benchmark into a single slide.

Produces:
  benchmark_summary.png   - 2x3 comparison figure (throughput + TTFT vs QPS)
  benchmark_summary.pptx  - one 16:9 slide: settings, chart, conclusions, table

Reuses the data loaders from gen_report.py so the numbers stay in sync with
report.md.
"""
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

import gen_report as gr

BASE = os.path.dirname(os.path.abspath(__file__))
PNG = os.path.join(BASE, "benchmark_summary.png")
PPTX = os.path.join(BASE, "benchmark_summary.pptx")

# baseline / cpu / fs -> colour (hex + RGB) + marker
STYLE = {
    "baseline":   dict(hex="#7f7f7f", rgb=(0x7f, 0x7f, 0x7f), marker="o",
                       label="Baseline"),
    "native-cpu": dict(hex="#1f77b4", rgb=(0x1f, 0x77, 0xb4), marker="s",
                       label="Native CPU offload"),
    "native-fs":  dict(hex="#2ca02c", rgb=(0x2c, 0xa0, 0x2c), marker="^",
                       label="Native FS offload"),
}
MODEL_SHORT = {
    "llama3-70b": "Llama-3 70B\n(TP=8 x4)",
    "qwen3-30b-a3b": "Qwen3-30B-A3B\n(TP=4 x8)",
    "qwen3-32b": "Qwen3-32B\n(TP=4 x8)",
}
MODEL_TBL = {
    "llama3-70b": "Llama-3 70B",
    "qwen3-30b-a3b": "Qwen3-30B-A3B",
    "qwen3-32b": "Qwen3-32B",
}
SETTINGS = [
    ("baseline", "Baseline (no offload)",
     "GPU-only prefix cache; evicts & recomputes once the reuse set exceeds VRAM."),
    ("native-cpu", "Native CPU offload",
     "KV blocks spill VRAM -> host RAM (OffloadingConnector); tier ~80% of reuse set."),
    ("native-fs", "Native FS offload",
     "KV blocks spill VRAM -> CPU -> node-local NVMe (TieringOffloadingSpec); holds ~full set."),
]


def series(case, key):
    xs, ys = [], []
    for q in sorted(case):
        v = case[q].get(key)
        if v is not None and case[q]["failures"] == 0:
            xs.append(q)
            ys.append(v)
    return xs, ys


def build_figure(data):
    fig, axes = plt.subplots(2, len(gr.MODELS), figsize=(12.8, 5.4), dpi=200)
    rows = [("out_tok_s", "Output throughput (tok/s)", False),
            ("ttft_mean", "TTFT mean (s)", True)]
    for r, (key, ylabel, logy) in enumerate(rows):
        for c, m in enumerate(gr.MODELS):
            ax = axes[r][c]
            for cfg in gr.CONFIGS:
                xs, ys = series(data[m][cfg], key)
                if not xs:
                    continue
                st = STYLE[cfg]
                ax.plot(xs, ys, marker=st["marker"], color=st["hex"],
                        label=st["label"], linewidth=1.8, markersize=4.5)
            if logy:
                ax.set_yscale("log")
            ax.grid(True, alpha=0.3, linewidth=0.5)
            if r == 0:
                ax.set_title(MODEL_SHORT[m], fontsize=11, fontweight="bold")
            if r == len(rows) - 1:
                ax.set_xlabel("Offered QPS (cluster)", fontsize=9)
            if c == 0:
                ax.set_ylabel(ylabel, fontsize=10)
            ax.tick_params(labelsize=8)
    handles, labels = axes[0][0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", ncol=3, fontsize=10,
               frameon=False, bbox_to_anchor=(0.5, 1.01))
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    fig.savefig(PNG, bbox_inches="tight")
    plt.close(fig)


def pct(a, b):
    return (a - b) / b * 100.0 if b else 0.0


def peak_qps(data, m):
    common = set(data[m][gr.CONFIGS[0]])
    for c in gr.CONFIGS[1:]:
        common &= set(data[m][c])
    clean = [q for q in common
             if all(data[m][c][q]["failures"] == 0 for c in gr.CONFIGS)]
    return max(clean) if clean else None


def build_conclusions(data, prefix):
    out = {}
    for m in gr.MODELS:
        q = peak_qps(data, m)
        base = data[m]["baseline"][q]
        fs = data[m]["native-fs"][q]
        _, ext_fs, _ = prefix[m]["native-fs"]
        out[m] = dict(q=q,
                      dt=pct(fs["out_tok_s"], base["out_tok_s"]),
                      dl=pct(fs["ttft_mean"], base["ttft_mean"]),
                      ext=ext_fs)
    return out


def build_table_rows(data, prefix):
    """Aggregated results at each model's peak sustained (failure-free) QPS."""
    header = ["Model", "Config", "Peak QPS", "Out tok/s",
              "TTFT mean (s)", "E2E p99 (s)", "Tier-hit %"]
    rows = [header]
    for m in gr.MODELS:
        q = peak_qps(data, m)
        for ci, c in enumerate(gr.CONFIGS):
            s = data[m][c][q]
            _, ext, _ = prefix[m][c]
            rows.append([
                MODEL_TBL[m] if ci == 0 else "",
                STYLE[c]["label"],
                f"{q:g}",
                f"{s['out_tok_s']:.0f}",
                f"{s['ttft_mean']:.2f}",
                f"{s['e2e_p99']:.1f}",
                "-" if ext is None else f"{ext:.0f}",
            ])
    return rows


def add_slide(concl, table_rows):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # ---- Title ----
    tb = slide.shapes.add_textbox(Inches(0.3), Inches(0.12),
                                  Inches(12.7), Inches(0.55))
    tf = tb.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = ("Tiered Prefix-Cache KV-Offloading on 32x Intel Arc Pro B60 - "
              "Baseline vs CPU vs FS")
    r.font.size = Pt(21)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

    # ---- 3-settings description (colour-keyed to chart) ----
    sb = slide.shapes.add_textbox(Inches(0.3), Inches(0.66),
                                  Inches(12.7), Inches(0.92))
    stf = sb.text_frame
    stf.word_wrap = True
    for i, (cfg, name, desc) in enumerate(SETTINGS):
        p = stf.paragraphs[0] if i == 0 else stf.add_paragraph()
        p.space_after = Pt(1)
        r1 = p.add_run()
        r1.text = f"{name}:  "
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(*STYLE[cfg]["rgb"])
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # ---- Chart (left) ----
    slide.shapes.add_picture(PNG, Inches(0.2), Inches(1.62), width=Inches(8.0))

    # ---- Conclusions (right) ----
    box = slide.shapes.add_textbox(Inches(8.4), Inches(1.6),
                                   Inches(4.8), Inches(3.45))
    tf = box.text_frame
    tf.word_wrap = True
    r0 = tf.paragraphs[0].add_run()
    r0.text = "Key findings"
    r0.font.size = Pt(14)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(0x0f, 0x0f, 0x0f)

    def head(text):
        p = tf.add_paragraph()
        p.space_before = Pt(4)
        rr = p.add_run()
        rr.text = text
        rr.font.size = Pt(11)
        rr.font.bold = True
        rr.font.color.rgb = RGBColor(0x15, 0x3a, 0x6b)

    def bullet(text, color=RGBColor(0x22, 0x22, 0x22)):
        p = tf.add_paragraph()
        p.space_after = Pt(1)
        rr = p.add_run()
        rr.text = "- " + text
        rr.font.size = Pt(9.5)
        rr.font.color.rgb = color

    cq = concl["qwen3-30b-a3b"]
    head("Qwen3-30B-A3B & Qwen3-32B - FS wins")
    bullet("FS best across the full 2-32 ladder; ~85-87% tier-hit.")
    bullet(f"At 32 QPS FS gives +{concl['qwen3-32b']['dt']:.0f}% tok/s / "
           f"{concl['qwen3-32b']['dl']:.0f}% TTFT (32B) and "
           f"+{cq['dt']:.0f}% / {cq['dl']:.0f}% (30B) vs baseline.")
    lc = concl["llama3-70b"]
    head("Llama-3 70B - load-dependent")
    bullet("Baseline best at light load (<=2 QPS): offload not yet amortized.")
    bullet(f"FS wins from ~4 QPS; +{lc['dt']:.0f}% tok/s / {lc['dl']:.0f}% "
           f"TTFT at {lc['q']:g} QPS, ~{lc['ext']:.0f}% tier-hit.")
    head("Bottom line")
    bullet("Native FS offload is the recommended tier for sustained load on "
           "all three models.", RGBColor(0x18, 0x60, 0x18))
    bullet("PCIe-bound (no Xe-Link) - pays off once recompute pressure "
           "exceeds transfer cost.")

    # ---- Aggregated data table (bottom, full width) ----
    cap = slide.shapes.add_textbox(Inches(0.3), Inches(5.08),
                                   Inches(12.7), Inches(0.3))
    rc = cap.text_frame.paragraphs[0].add_run()
    rc.text = "Aggregated results at each model's peak sustained (failure-free) QPS"
    rc.font.size = Pt(11)
    rc.font.bold = True
    rc.font.color.rgb = RGBColor(0x0f, 0x0f, 0x0f)

    nrows = len(table_rows)
    ncols = len(table_rows[0])
    gtbl = slide.shapes.add_table(nrows, ncols, Inches(0.3), Inches(5.4),
                                  Inches(12.73), Inches(1.9)).table
    gtbl.columns[0].width = Inches(2.2)
    gtbl.columns[1].width = Inches(2.5)
    for ci in range(2, ncols):
        gtbl.columns[ci].width = Inches(1.606)
    for ri, row in enumerate(table_rows):
        for ci, val in enumerate(row):
            cell = gtbl.cell(ri, ci)
            cell.margin_top = Pt(1)
            cell.margin_bottom = Pt(1)
            cell.margin_left = Pt(4)
            cell.margin_right = Pt(4)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            run = para.add_run()
            run.text = val
            run.font.size = Pt(9)
            if ri == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x2f, 0x4b, 0x6e)
            else:
                cfg = gr.CONFIGS[(ri - 1) % 3]
                if ci == 1:
                    run.font.color.rgb = RGBColor(*STYLE[cfg]["rgb"])
                    run.font.bold = True
                else:
                    run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
                cell.fill.solid()
                band = ((ri - 1) // 3) % 2
                cell.fill.fore_color.rgb = (RGBColor(0xf2, 0xf5, 0xf9) if band
                                            else RGBColor(0xff, 0xff, 0xff))

    prs.save(PPTX)


def main():
    data = {m: {c: gr.load_case(m, c) for c in gr.CONFIGS} for m in gr.MODELS}
    prefix = {m: {c: gr.load_prefix(m, c) for c in gr.CONFIGS} for m in gr.MODELS}
    build_figure(data)
    concl = build_conclusions(data, prefix)
    table_rows = build_table_rows(data, prefix)
    add_slide(concl, table_rows)
    print("wrote", PNG)
    print("wrote", PPTX)
    for m in gr.MODELS:
        c = concl[m]
        print(f"  {m}: @QPS {c['q']:g}  FS vs base  out+{c['dt']:.1f}%  "
              f"ttft{c['dl']:.1f}%  ext-hit {c['ext']:.0f}%")


if __name__ == "__main__":
    main()
